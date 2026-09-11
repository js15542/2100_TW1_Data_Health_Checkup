from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("02_TW1_Part1_Handoff.docx")
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"
DARK_GRAY = "595959"
BLACK = RGBColor(0, 0, 0)


def set_font(run, name="Arial", size=10.5, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths=None, font_size=9.0, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])

    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell, 80, 90, 80, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(header), size=9.2, bold=True, color=RGBColor(255, 255, 255))

    for row_idx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col, value in enumerate(row_data):
            cell = row.cells[col]
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if alignments and col < len(alignments):
                p.alignment = alignments[col]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)), size=font_size)

    if widths:
        for row in table.rows:
            for col, width in enumerate(widths):
                row.cells[col].width = Inches(width)

    table.rows[-1]._tr.addnext(OxmlElement("w:bookmarkStart")) if False else None
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), size=14 if level == 1 else 11.5, bold=True)
    return p


def add_body(doc, text, bold_prefix=None, italic=False, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.04
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), bold=True)
        set_font(p.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_font(p.add_run(text), italic=italic)
    return p


def add_bullet(doc, label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(label), size=9.9, bold=True)
    set_font(p.add_run(text), size=9.9)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10.5)
for style_name in ("Title", "Heading 1", "Heading 2"):
    styles[style_name].font.color.rgb = BLACK
    styles[style_name].font.name = "Arial"
    styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(4)
set_font(title.add_run("TW1 Part 1 Handoff"), size=21, bold=True)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(5)
set_font(meta.add_run("Source  "), size=9, bold=True, color=RGBColor(89, 89, 89))
set_font(meta.add_run("B2B SaaS Churn Data.xlsx    "), size=9, color=RGBColor(89, 89, 89))
set_font(meta.add_run("Checked  "), size=9, bold=True, color=RGBColor(89, 89, 89))
set_font(meta.add_run("September 10 2026"), size=9, color=RGBColor(89, 89, 89))

add_body(
    doc,
    "This handoff provides the six-part structure for the final one-page Red Flags and Strategy Memo. Sections 1 and 5 are complete; teammates should fill Sections 2, 3, 4, and 6 using the verified evidence that follows.",
    after=5,
)

add_heading(doc, "Deliverable Template", 1)

add_heading(doc, "1 Data Quality Grade Completed", 2)
add_body(doc, "Grade C", bold_prefix="Grade C", after=2)
add_body(
    doc,
    "The extract is readable and contains 2,500 unique customers, but nine exact duplicates, one conflicting customer ID, inconsistent churn and industry labels, impossible values, and two extreme outliers would distort analysis and campaign targeting. The data can support a high-level churn summary after documented cleaning, but Last_Login_Days_Ago requires a common snapshot date before it can be used reliably in predictive modeling.",
    after=4,
)

add_heading(doc, "2 Top 3 Critical Errors and Strategic Business Impact", 2)
add_body(
    doc,
    "Team to complete. Select three issues from Part 1 Audit Evidence and explain how each would skew the retention campaign.",
    italic=True,
    after=3,
)

add_heading(doc, "3 Data Leakage Warning", 2)
add_body(
    doc,
    "Team to complete. Identify the column to remove before predictive modeling and explain why it reveals the outcome.",
    italic=True,
    after=3,
)

add_heading(doc, "4 High Level Churn Observation", 2)
add_body(
    doc,
    "Team to complete. Summarize the direction of the 2022 to 2023 change using the Clean results in Part 1 Audit Evidence.",
    italic=True,
    after=3,
)

add_heading(doc, "5 Data Cleaning Strategy Completed", 2)
add_bullet(doc, "Duplicates and labels. ", "Remove the nine exact duplicates. For Customer_ID 1002, the reproducible Clean view keeps the first record, coded as 1, and temporarily maps it to No because Date_Cancelled is blank. Keep the conflict flagged for source confirmation.")
add_bullet(doc, "Categories. ", "Trim spaces and standardize Industry capitalization, reducing 12 spellings to six industries. Keep Contract_Type Unknown as a separate category with a flag.")
add_bullet(doc, "Missing support tickets. ", "Fill blank Support_Tickets values with 0 and keep a missing-value flag. This produces a Clean mean of 1.88. If the team chooses median imputation, use 2; the resulting mean is 1.98.")
add_bullet(doc, "Logical errors. ", "Map Churn_Label 1 to Yes only when Date_Cancelled is present; otherwise map it to No, and retain a reconciliation flag. Set negative user, revenue, and login-day values to missing and add flags. Do not take absolute values or guess corrections.")
add_bullet(doc, "Extreme outliers. ", "Set the $1,000,000 revenue value and 50,000-user value to missing for baseline and modeling calculations, retain review flags, and confirm them with the source before correction or deletion.")
add_bullet(doc, "Modeling fields. ", "Remove Date_Cancelled and Company_Name from model features. Do not use Last_Login_Days_Ago until IT provides a common snapshot date or re-extracts the field using one reference date.")

add_heading(doc, "6 Strategic Data Augmentation Recommendations", 2)
add_body(
    doc,
    "Team to complete. Recommend two or three relevant Zero Party, First Party, or Third Party attributes and explain how each would improve the retention campaign.",
    italic=True,
    after=0,
)

doc.add_page_break()

add_heading(doc, "Part 1 Audit Evidence", 1)
add_body(
    doc,
    "The tables below contain the verified Raw and Clean figures teammates need for the remaining memo sections. No further calculation is required. Use Clean results for business conclusions and Raw comparisons for data-quality and outlier impact statements.",
    after=5,
)

add_heading(doc, "Volume and Redundancy", 2)
add_table(
    doc,
    ["Check", "Result"],
    [
        ("Rows in the delivered file", "2,510; IT claimed 2,500"),
        ("Unique Customer_ID values", "2,500"),
        ("Exact duplicate rows", "9"),
        ("Rows after exact duplicates", "2,501"),
        ("Additional conflicting Customer_ID", "1; Customer_ID 1002 has No versus 1"),
        ("Final Clean customers", "2,500"),
    ],
    widths=[2.55, 4.45],
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
)

add_heading(doc, "Schema and Format", 2)
add_table(
    doc,
    ["Check", "Result"],
    [
        ("Numeric storage", "All numeric fields are stored as numbers; no currency strings were found"),
        ("Monthly_Revenue display", "378 cells use currency format; 2,132 use General format"),
        ("Industry", "12 raw spellings for six industries; 251 lower-case rows in Raw and 250 after exact-duplicate removal"),
        ("Churn_Label", "1,974 No; 486 Yes; 50 numeric 1 entries"),
        ("Numeric 1 reconciliation", "7 have Date_Cancelled; 43 do not"),
        ("Contract_Type", "50 Unknown entries"),
        ("Support_Tickets", "125 blanks, equal to 5.0% of Raw rows"),
        ("Date_Cancelled", "493 populated; 2,017 blank, normally expected for retained customers"),
    ],
    widths=[2.15, 4.85],
)

add_heading(doc, "Logical Errors", 2)
add_table(
    doc,
    ["Field", "Problem", "Count"],
    [
        ("Total_Users", "-10; Customer_ID 1062", "1"),
        ("Monthly_Revenue", "-$500; Customer_ID 1052", "1"),
        ("Last_Login_Days_Ago", "-5", "50"),
        ("Churn_Label", "Coded as 1 without Date_Cancelled", "43"),
        ("Date_Cancelled", "Earlier than Join_Date", "0"),
        ("Churn_Label No", "Date_Cancelled populated", "0"),
    ],
    widths=[2.0, 4.2, 0.8],
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
)
add_body(doc, "The 43 rows coded as 1 without a cancellation date are label inconsistencies, not confirmed churn.", italic=True, after=4)

add_heading(doc, "Label Reconciliation", 2)
add_table(
    doc,
    ["Stage", "Yes", "No", "Ambiguous 1", "Total"],
    [
        ("Raw file", "486", "1,974", "50", "2,510"),
        ("After exact duplicates", "484", "1,967", "50", "2,501"),
        ("After Customer_ID 1002 resolution", "484", "1,966", "50", "2,500"),
        ("Clean view after date-based mapping", "491", "2,009", "0", "2,500"),
    ],
    widths=[3.0, 0.8, 0.9, 1.3, 1.0],
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    font_size=8.8,
)

add_heading(doc, "Final Clean Metrics", 2)
add_table(
    doc,
    ["Metric", "Numerator and Denominator", "Clean Result"],
    [
        ("Overall churn", "491 / 2,500", "19.6%"),
        ("2022 churn", "145 / 1,653 active customers", "8.8%"),
        ("2023 churn", "346 / 2,355 active customers", "14.7%"),
        ("Change in cancellations", "346 / 145", "2.39x"),
        ("Largest industry", "Retail 445 / 2,500", "17.8% of customers"),
        ("Highest industry churn", "Education 91 / 401", "22.7%"),
        ("Month-to-Month churn", "268 / 817", "32.8%"),
        ("1 Year churn", "107 / 820", "13.0%"),
        ("2 Year churn", "108 / 813", "13.3%"),
        ("Unknown contract churn", "8 / 50", "16.0%"),
    ],
    widths=[2.3, 3.1, 1.6],
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
)

add_heading(doc, "Extreme Outliers and Baseline Impact", 2)
add_table(
    doc,
    ["Field", "Outlier", "Raw Effect", "Clean Baseline"],
    [
        ("Monthly_Revenue", "$1,000,000; Customer_ID 1051", "$1,381.36 mean with it; $983.35 without it; +$398.01 or +40.5%", "$984.97 mean after all cleaning rules"),
        ("Total_Users", "50,000; Customer_ID 1061", "39.28 mean with it; 19.37 without it; +19.91 or +102.8%", "19.40 mean after all cleaning rules"),
    ],
    widths=[1.35, 1.85, 2.65, 1.65],
    font_size=8.5,
)

add_heading(doc, "Missing Values and Engagement", 2)
add_table(
    doc,
    ["Check", "Raw", "Clean or Interpretation"],
    [
        ("Support_Tickets", "125 blanks; non-missing mean 1.98", "0 imputation plus missing flag; mean 1.88. Median 2 is the documented alternative and would produce mean 1.98"),
        ("Last_Login_Days_Ago", "50 rows at -5; mean 14.41", "Negative values set to missing; mean 14.82, but the field remains unusable without a common snapshot date"),
        ("Login after cancellation check", "470 of 493 cancelled rows imply a later login", "462 of 491 imply a later login; six Clean values are missing after the negative-value rule"),
    ],
    widths=[1.8, 2.25, 3.45],
    font_size=8.6,
)

add_heading(doc, "Additional Reliability Evidence", 2)
p = add_bullet(doc, "Cancellation history. ", "Join_Date spans 2020 through 2022, but Date_Cancelled contains only 2022 and 2023 exits. Confirm whether earlier churn was removed before making a longer-term trend claim.")
p.paragraph_format.keep_with_next = True
p = add_bullet(doc, "Leakage evidence. ", "Date_Cancelled is populated for all 486 Raw rows labeled Yes and none of the 1,974 rows labeled No. Among the 50 numeric 1 rows, seven have a date and 43 do not.")
p.paragraph_format.keep_with_next = True
p = add_bullet(doc, "Identifier evidence. ", "Company_Name is derived from Customer_ID for every row and should not be used as a predictive feature.")
p.paragraph_format.keep_with_next = True
add_bullet(doc, "Pricing sanity check. ", "Typical accounts pay about $50.80 per user per month. The revenue whale implies about $58,824 per user, while the 50,000-user row implies about $0.02 per user, supporting source review rather than guessed corrections.")

add_heading(doc, "Reference Files", 2)
add_body(doc, "TW1_Audit_Findings.xlsx contains the offending rows for each audit check. TW1_Raw_vs_Clean_Metrics.xlsx contains the Raw and Clean comparisons, including the Metrics, Label Reconciliation, segment, outlier, and leakage tables.", after=0)

doc.core_properties.title = "TW1 Part 1 Handoff"
doc.core_properties.subject = "Data Quality Grade, Data Cleaning Strategy, and audit evidence"
doc.core_properties.author = "Team session 2-4 5"
doc.core_properties.keywords = "TW1, data health audit, churn, cleaning strategy"
doc.save(OUTPUT)
print(OUTPUT)
